"""
审批清单 Word 导出路由
优先接收结构化清单数据并导出可编辑 DOCX；兼容旧版图片导出请求。
"""
from __future__ import annotations

import base64
import concurrent.futures
import io
import logging
import os
import re
import threading
import uuid
from datetime import datetime, timedelta, timezone
from typing import Final
from urllib.parse import quote

from fastapi import APIRouter, HTTPException, Response
from PIL import Image
from pydantic import BaseModel, Field

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Emu, RGBColor
    from docx.oxml.ns import qn
    DOCX_IMPORT_ERROR: ModuleNotFoundError | None = None
except ModuleNotFoundError as exc:  # pragma: no cover - runtime guard
    Document = None  # type: ignore[assignment]
    WD_ALIGN_PARAGRAPH = None  # type: ignore[assignment]
    Emu = None  # type: ignore[assignment]
    RGBColor = None  # type: ignore[assignment]
    qn = None  # type: ignore[assignment]
    DOCX_IMPORT_ERROR = exc

router = APIRouter()
logger = logging.getLogger(__name__)

DATA_URL_PATTERN: Final[re.Pattern[str]] = re.compile(
    r"^data:image/(?P<format>[a-zA-Z0-9.+-]+);base64,(?P<data>.+)$"
)
DEFAULT_FILE_NAME = "审核清单.docx"
ASCII_FALLBACK_FILE_NAME = "approval-checklist.docx"
DEFAULT_FONT_NAME = "微软雅黑"
COLOR_TEXT_DEFAULT = RGBColor(0, 0, 0)
COLOR_TEXT_PASS = RGBColor(22, 163, 74)
COLOR_TEXT_FAIL = RGBColor(220, 38, 38)


def _read_int_env(name: str, default: int, *, min_value: int, max_value: int) -> int:
    raw_value = os.getenv(name)
    if raw_value is None:
        return default
    try:
        parsed = int(raw_value)
    except ValueError:
        logger.warning("%s 不是有效整数，使用默认值 %s", name, default)
        return default
    return max(min_value, min(max_value, parsed))


JOB_RETENTION_MINUTES = _read_int_env(
    "HDMS_EXPORT_JOB_RETENTION_MINUTES",
    default=30,
    min_value=5,
    max_value=12 * 60,
)
JOB_RUNNING_TIMEOUT_MINUTES = _read_int_env(
    "HDMS_EXPORT_JOB_RUNNING_TIMEOUT_MINUTES",
    default=120,
    min_value=10,
    max_value=24 * 60,
)
JOB_MAX_WORKERS = _read_int_env(
    "HDMS_EXPORT_JOB_MAX_WORKERS",
    default=2,
    min_value=1,
    max_value=4,
)


class ExportWordRequest(BaseModel):
    project_name: str | None = Field(default=None, max_length=100)
    exported_date: str | None = Field(default=None, max_length=32)
    items: list["ExportWordItem"] = Field(default_factory=list)
    image_data_url: str | None = Field(default=None, min_length=20)
    file_name: str | None = Field(default=None, max_length=120)


class PassedPlot(BaseModel):
    name: str = ""
    buildings: list[str] = Field(default_factory=list)


class FailedIssue(BaseModel):
    plotName: str = ""
    buildingName: str = ""
    issue: str = ""
    details: str | None = None


class DetailedPassed(BaseModel):
    plots: list[PassedPlot] = Field(default_factory=list)
    totalPlots: int = 0
    totalBuildings: int = 0


class DetailedFailed(BaseModel):
    items: list[FailedIssue] = Field(default_factory=list)
    totalPlots: int = 0
    totalBuildings: int = 0


class DetailedStats(BaseModel):
    passed: DetailedPassed = Field(default_factory=DetailedPassed)
    failed: DetailedFailed = Field(default_factory=DetailedFailed)
    summary: str = ""


class ExportWordItem(BaseModel):
    id: str
    name: str = ""
    is_pass: bool = False
    summary: str = ""
    detailed_stats: DetailedStats | None = None
    gov_suggestion: str = ""


ExportWordRequest.model_rebuild()


class ExportJobCreateResponse(BaseModel):
    job_id: str
    status: str


class ExportJobStatusResponse(BaseModel):
    job_id: str
    status: str
    file_name: str
    created_at: str
    updated_at: str
    completed_at: str | None = None
    error: str | None = None
    download_ready: bool = False


class _ExportJobRecord:
    def __init__(self, job_id: str, file_name: str):
        now = datetime.now(timezone.utc)
        self.job_id = job_id
        self.status = "queued"
        self.file_name = file_name
        self.created_at = now
        self.updated_at = now
        self.completed_at: datetime | None = None
        self.error: str | None = None
        self.word_bytes: bytes | None = None

    def to_response(self) -> ExportJobStatusResponse:
        return ExportJobStatusResponse(
            job_id=self.job_id,
            status=self.status,
            file_name=self.file_name,
            created_at=self.created_at.isoformat(),
            updated_at=self.updated_at.isoformat(),
            completed_at=self.completed_at.isoformat() if self.completed_at else None,
            error=self.error,
            download_ready=self.word_bytes is not None and self.status == "completed",
        )


class ExportJobManager:
    def __init__(self):
        self._jobs: dict[str, _ExportJobRecord] = {}
        self._lock = threading.Lock()
        self._executor = concurrent.futures.ThreadPoolExecutor(
            max_workers=JOB_MAX_WORKERS,
            thread_name_prefix="word-export",
        )

    def create_job(self, request: ExportWordRequest) -> _ExportJobRecord:
        request_copy = request.model_copy(deep=True)
        job_id = uuid.uuid4().hex
        record = _ExportJobRecord(job_id=job_id, file_name=_resolve_file_name(request_copy))
        with self._lock:
            self._purge_expired_locked()
            self._jobs[job_id] = record
        self._executor.submit(self._run_job, job_id, request_copy)
        return record

    def get_job(self, job_id: str) -> _ExportJobRecord:
        with self._lock:
            self._purge_expired_locked()
            record = self._jobs.get(job_id)
            if record is None:
                raise HTTPException(status_code=404, detail="导出任务不存在或已过期")
            return record

    def _run_job(self, job_id: str, request: ExportWordRequest) -> None:
        self._update_status(job_id, status="running")
        try:
            if request.items:
                word_bytes = _build_word_from_items(request)
            elif request.image_data_url:
                image_bytes = _parse_image_data_url(request.image_data_url)
                word_bytes = _build_word_bytes(image_bytes)
            else:
                raise HTTPException(status_code=400, detail="导出参数无效：items 或 image_data_url 至少提供一个")
        except HTTPException as exc:
            self._mark_failed(job_id, exc.detail if isinstance(exc.detail, str) else str(exc.detail))
            return
        except Exception as exc:  # pragma: no cover - defensive fallback
            logger.warning("Word 异步导出任务失败: %s (%s)", job_id, exc)
            self._mark_failed(job_id, str(exc))
            return

        self._mark_completed(job_id, word_bytes)

    def _update_status(self, job_id: str, status: str) -> None:
        with self._lock:
            record = self._jobs.get(job_id)
            if record is None:
                return
            record.status = status
            record.updated_at = datetime.now(timezone.utc)
            if status == "running":
                record.error = None

    def _mark_completed(self, job_id: str, word_bytes: bytes) -> None:
        with self._lock:
            record = self._jobs.get(job_id)
            if record is None:
                return
            record.status = "completed"
            record.word_bytes = word_bytes
            record.error = None
            record.updated_at = datetime.now(timezone.utc)
            record.completed_at = record.updated_at

    def _mark_failed(self, job_id: str, error: str) -> None:
        with self._lock:
            record = self._jobs.get(job_id)
            if record is None:
                return
            record.status = "failed"
            record.error = error
            record.word_bytes = None
            record.updated_at = datetime.now(timezone.utc)
            record.completed_at = record.updated_at

    def _purge_expired_locked(self) -> None:
        now = datetime.now(timezone.utc)
        finished_deadline = now - timedelta(minutes=JOB_RETENTION_MINUTES)
        running_deadline = now - timedelta(minutes=JOB_RUNNING_TIMEOUT_MINUTES)

        expired_job_ids: list[str] = []
        for job_id, record in self._jobs.items():
            if record.status in {"completed", "failed"} and record.updated_at < finished_deadline:
                expired_job_ids.append(job_id)
            elif record.status in {"queued", "running"} and record.updated_at < running_deadline:
                expired_job_ids.append(job_id)

        for job_id in expired_job_ids:
            del self._jobs[job_id]


JOB_MANAGER = ExportJobManager()


def _ensure_docx_available() -> None:
    if DOCX_IMPORT_ERROR is None:
        return
    raise HTTPException(
        status_code=503,
        detail="Word 导出依赖缺失：python-docx 未安装，请在 approval_checklist 环境执行 pip install python-docx",
    ) from DOCX_IMPORT_ERROR


def _parse_image_data_url(image_data_url: str) -> bytes:
    match = DATA_URL_PATTERN.match(image_data_url.strip())
    if not match:
        raise HTTPException(status_code=400, detail="image_data_url 格式无效，需为 data:image/...;base64,...")
    try:
        return base64.b64decode(match.group("data"), validate=True)
    except Exception as exc:  # pragma: no cover - defensive branch
        raise HTTPException(status_code=400, detail="image_data_url Base64 解码失败") from exc


def _normalize_image(image: Image.Image) -> Image.Image:
    if image.mode in ("RGBA", "LA"):
        background = Image.new("RGB", image.size, (255, 255, 255))
        alpha = image.getchannel("A") if "A" in image.getbands() else None
        background.paste(image.convert("RGB"), mask=alpha)
        return background
    if image.mode != "RGB":
        return image.convert("RGB")
    return image


def _build_word_bytes(image_bytes: bytes) -> bytes:
    _ensure_docx_available()
    with Image.open(io.BytesIO(image_bytes)) as source_image:
        normalized = _normalize_image(source_image)
        document = Document()
        section = document.sections[0]

        content_width = int(section.page_width - section.left_margin - section.right_margin)
        content_height = int(section.page_height - section.top_margin - section.bottom_margin)
        if content_width <= 0:
            content_width = int(6.0 * 914400)  # 6 inch fallback
        if content_height <= 0:
            content_height = int(9.0 * 914400)  # 9 inch fallback

        px_per_emu = normalized.width / content_width if content_width > 0 else 1
        max_slice_height = max(1, int(content_height * px_per_emu))

        offset = 0
        while offset < normalized.height:
            slice_height = min(max_slice_height, normalized.height - offset)
            image_slice = normalized.crop((0, offset, normalized.width, offset + slice_height))

            slice_buffer = io.BytesIO()
            image_slice.save(slice_buffer, format="PNG")
            slice_buffer.seek(0)

            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(slice_buffer, width=Emu(content_width))

            offset += slice_height
            if offset < normalized.height:
                document.add_page_break()

        output = io.BytesIO()
        document.save(output)
        output.seek(0)
        return output.getvalue()


def _set_run_font(run, color: RGBColor = COLOR_TEXT_DEFAULT) -> None:
    run.font.name = DEFAULT_FONT_NAME
    run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), DEFAULT_FONT_NAME)
    run._element.rPr.rFonts.set(qn("w:ascii"), DEFAULT_FONT_NAME)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), DEFAULT_FONT_NAME)


def _set_paragraph_text(paragraph, value: str, color: RGBColor = COLOR_TEXT_DEFAULT) -> None:
    paragraph.clear()
    run = paragraph.add_run(value)
    if color != COLOR_TEXT_DEFAULT:
        run.font.color.rgb = color
    paragraph.paragraph_format.space_before = 0
    paragraph.paragraph_format.space_after = 0


def _set_table_cell_text(cell, value: str, color: RGBColor = COLOR_TEXT_DEFAULT) -> None:
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    _set_paragraph_text(paragraph, value, color=color)


def _apply_document_default_font(document, font_name: str = DEFAULT_FONT_NAME) -> None:
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        try:
            style = document.styles[style_name]
        except KeyError:
            continue
        style.font.name = font_name
        if style.element.rPr is None:
            style.element.get_or_add_rPr()
        style.element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        style.element.rPr.rFonts.set(qn("w:ascii"), font_name)
        style.element.rPr.rFonts.set(qn("w:hAnsi"), font_name)


def _build_word_from_items(request: ExportWordRequest) -> bytes:
    _ensure_docx_available()
    document = Document()
    _apply_document_default_font(document)
    project_name = (request.project_name or "审核清单").strip()
    report_date = request.exported_date or ""

    title = document.add_heading(f"{project_name}管控审核清单", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        _set_run_font(run)
    if report_date:
        date_paragraph = document.add_paragraph(f"日期：{report_date}")
        date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in date_paragraph.runs:
            _set_run_font(run)

    if request.items:
        document.add_heading("检测项目录", level=1)
        catalog_table = document.add_table(rows=1, cols=3)
        catalog_table.style = "Table Grid"
        _set_table_cell_text(catalog_table.rows[0].cells[0], "序号")
        _set_table_cell_text(catalog_table.rows[0].cells[1], "检测项")
        _set_table_cell_text(catalog_table.rows[0].cells[2], "检测结果")

        for index, item in enumerate(request.items, start=1):
            row_cells = catalog_table.add_row().cells
            _set_table_cell_text(row_cells[0], str(index))
            _set_table_cell_text(row_cells[1], item.name)
            summary_color = COLOR_TEXT_PASS if item.is_pass else COLOR_TEXT_FAIL
            _set_table_cell_text(row_cells[2], item.summary, color=summary_color)

    for index, item in enumerate(request.items, start=1):
        document.add_page_break()
        document.add_heading(f"{index}. {item.name}", level=1)
        result_paragraph = document.add_paragraph()
        _set_paragraph_text(result_paragraph, "检测结果：")
        result_run = result_paragraph.add_run(item.summary)
        _set_run_font(result_run, color=COLOR_TEXT_PASS if item.is_pass else COLOR_TEXT_FAIL)

        stats = item.detailed_stats
        if stats:
            if stats.summary:
                summary_paragraph = document.add_paragraph()
                _set_paragraph_text(summary_paragraph, f"统计摘要：{stats.summary}")

            if stats.passed.plots:
                passed_title = document.add_paragraph()
                _set_paragraph_text(
                    passed_title,
                    f"通过项（{stats.passed.totalBuildings} 栋）",
                    color=COLOR_TEXT_PASS,
                )
                passed_table = document.add_table(rows=1, cols=3)
                passed_table.style = "Table Grid"
                _set_table_cell_text(passed_table.rows[0].cells[0], "序号")
                _set_table_cell_text(passed_table.rows[0].cells[1], "地块")
                _set_table_cell_text(passed_table.rows[0].cells[2], "建筑")
                for row_index, plot in enumerate(stats.passed.plots, start=1):
                    row_cells = passed_table.add_row().cells
                    _set_table_cell_text(row_cells[0], str(row_index))
                    _set_table_cell_text(row_cells[1], plot.name)
                    _set_table_cell_text(row_cells[2], "、".join(plot.buildings) if plot.buildings else "-")

            if stats.failed.items:
                failed_title = document.add_paragraph()
                _set_paragraph_text(
                    failed_title,
                    f"不通过项（{stats.failed.totalBuildings} 栋）",
                    color=COLOR_TEXT_FAIL,
                )
                failed_table = document.add_table(rows=1, cols=5)
                failed_table.style = "Table Grid"
                headers = ["序号", "地块", "建筑/对象", "问题", "详情"]
                for col, header in enumerate(headers):
                    _set_table_cell_text(failed_table.rows[0].cells[col], header)
                for row_index, failed in enumerate(stats.failed.items, start=1):
                    row_cells = failed_table.add_row().cells
                    _set_table_cell_text(row_cells[0], str(row_index))
                    _set_table_cell_text(row_cells[1], failed.plotName)
                    _set_table_cell_text(row_cells[2], failed.buildingName)
                    _set_table_cell_text(row_cells[3], failed.issue)
                    _set_table_cell_text(row_cells[4], failed.details or "-")

        suggestion_title = document.add_paragraph()
        _set_paragraph_text(suggestion_title, "审查方建议：")
        suggestion_body = document.add_paragraph()
        _set_paragraph_text(suggestion_body, item.gov_suggestion.strip() or "（未填写）")

    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    return output.getvalue()


def _resolve_file_name(request: ExportWordRequest) -> str:
    if request.file_name:
        return request.file_name if request.file_name.lower().endswith(".docx") else f"{request.file_name}.docx"
    if request.project_name:
        return f"{request.project_name}_审核清单.docx"
    return DEFAULT_FILE_NAME


@router.post("/export-word")
async def export_word(request: ExportWordRequest) -> Response:
    _ensure_docx_available()

    try:
        if request.items:
            word_bytes = _build_word_from_items(request)
        elif request.image_data_url:
            image_bytes = _parse_image_data_url(request.image_data_url)
            word_bytes = _build_word_bytes(image_bytes)
        else:
            raise HTTPException(status_code=400, detail="导出参数无效：items 或 image_data_url 至少提供一个")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Word 生成失败: {exc}") from exc

    file_name = _resolve_file_name(request)
    content_disposition = (
        f"attachment; filename={ASCII_FALLBACK_FILE_NAME}; "
        f"filename*=UTF-8''{quote(file_name)}"
    )
    return Response(
        content=word_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": content_disposition},
    )


@router.post("/export-word/jobs", response_model=ExportJobCreateResponse, status_code=202)
async def create_export_word_job(request: ExportWordRequest) -> ExportJobCreateResponse:
    _ensure_docx_available()
    if not request.items and not request.image_data_url:
        raise HTTPException(status_code=400, detail="导出参数无效：items 或 image_data_url 至少提供一个")

    record = JOB_MANAGER.create_job(request)
    return ExportJobCreateResponse(job_id=record.job_id, status=record.status)


@router.get("/export-word/jobs/{job_id}", response_model=ExportJobStatusResponse)
async def get_export_word_job_status(job_id: str) -> ExportJobStatusResponse:
    record = JOB_MANAGER.get_job(job_id)
    return record.to_response()


@router.get("/export-word/jobs/{job_id}/download")
async def download_export_word_job(job_id: str) -> Response:
    record = JOB_MANAGER.get_job(job_id)
    if record.status in {"queued", "running"}:
        raise HTTPException(status_code=409, detail="导出任务仍在处理中")
    if record.status == "failed":
        raise HTTPException(status_code=409, detail=record.error or "导出任务失败")
    if not record.word_bytes:
        raise HTTPException(status_code=500, detail="导出任务结果异常")

    content_disposition = (
        f"attachment; filename={ASCII_FALLBACK_FILE_NAME}; "
        f"filename*=UTF-8''{quote(record.file_name)}"
    )
    return Response(
        content=record.word_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": content_disposition},
    )
